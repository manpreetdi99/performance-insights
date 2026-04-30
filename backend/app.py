from fastapi import FastAPI, HTTPException, Query
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
from db import get_connection, get_available_databases
import time
import os
import openpyxl

# ---- Antennas cache ----
_antennas_cache: list[dict] | None = None

ANTENNAS_EXCEL_PATH = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
    "cosmote all antennas",
    "geo4g.xlsx"
)

def _load_antennas() -> list[dict]:
    global _antennas_cache
    if _antennas_cache is not None:
        return _antennas_cache

    wb = openpyxl.load_workbook(ANTENNAS_EXCEL_PATH, read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(min_row=1, values_only=True)
    header = next(rows_iter)
    # Column indices (0-based)
    lat_idx      = header.index("Latitude_Sector")
    lon_idx      = header.index("Longitude_Sector")
    site_idx     = header.index("SiteID")
    cell_idx     = header.index("Cell_ID")
    name_idx     = header.index("CellName")
    az_idx       = header.index("Azimuth")
    freq_idx     = header.index("FREQUENCY")
    vendor_idx   = header.index("VENDOR")
    enb_idx      = header.index("ENB_NAME")
    tech_idx     = header.index("Technology")
    status_idx   = header.index("TECHSTATUS")
    pci_idx      = header.index("PCI")
    tilt_idx     = header.index("Downtilt")
    ht_idx       = header.index("HT")

    result = []
    for row in rows_iter:
        lat = row[lat_idx]
        lon = row[lon_idx]
        if lat is None or lon is None:
            continue
        try:
            lat = float(lat)
            lon = float(lon)
        except (TypeError, ValueError):
            continue
        result.append({
            "lat": lat,
            "lon": lon,
            "siteId":   row[site_idx],
            "cellId":   row[cell_idx],
            "cellName": row[name_idx],
            "azimuth":  row[az_idx],
            "freq":     row[freq_idx],
            "vendor":   row[vendor_idx],
            "enbName":  row[enb_idx],
            "tech":     row[tech_idx],
            "status":   row[status_idx],
            "pci":      row[pci_idx],
            "downtilt": row[tilt_idx],
            "height":   row[ht_idx],
        })
    wb.close()
    _antennas_cache = result
    return result

app = FastAPI()

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],  # μετά μπορείς να το περιορίσεις
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

from typing import Union

class QueryRequest(BaseModel):
    database: str
    queries: list[str]

class CommentRequest(BaseModel):
    database: str
    session_id: Union[str, int]
    comment: str | None = ""

@app.post("/api/calls/comment")
def update_call_comment(req: CommentRequest):
    try:
        conn = get_connection(req.database)
        cursor = conn.cursor()

        # Check if comment already exists in AnalysisComment
        cursor.execute("SELECT commentID FROM AnalysisComment WHERE Comment = ?", (req.comment,))
        row = cursor.fetchone()
        
        if row:
            comment_id = row[0]
        else:
            try:
                # We must insert it. In SQL Server OUTPUT INSERTED is supported.
                cursor.execute("INSERT INTO AnalysisComment (Comment) OUTPUT INSERTED.commentID VALUES (?)", (req.comment,))
                comment_id = cursor.fetchone()[0]
            except Exception as e:
                # Fallback if OUTPUT INSERTED is not supported or identity fails
                cursor.execute("INSERT INTO AnalysisComment (Comment) VALUES (?)", (req.comment,))
                cursor.execute("SELECT @@IDENTITY")
                comment_id = cursor.fetchone()[0]

        # check if it exists in bridge
        cursor.execute("SELECT sessionID FROM AnalysisCommentSessionsBridge WHERE sessionID = ?", (req.session_id,))
        if cursor.fetchone():
            cursor.execute("UPDATE AnalysisCommentSessionsBridge SET commentId = ? WHERE sessionID = ?", (comment_id, req.session_id))
        else:
            cursor.execute("INSERT INTO AnalysisCommentSessionsBridge (sessionID, commentId) VALUES (?, ?)", (req.session_id, comment_id))
            
        # If comment starts with 'fake' or 'FAKE', set session as invalid (Valid = 0), otherwise Valid = 1
        if req.comment and req.comment.lower().startswith("fake"):
            cursor.execute("UPDATE Sessions SET Valid = 0 WHERE SessionId = ?", (req.session_id,))
        else:
            cursor.execute("UPDATE Sessions SET Valid = 1 WHERE SessionId = ?", (req.session_id,))

        conn.commit()
        conn.close()

        return {"message": "Comment updated successfully"}
    except Exception as e:
        print(f"Error in update_call_comment bridge update: {e}")
        # If the above fails because of missing table, fallback to updating Sessions.InvalidReason
        try:
            conn = get_connection(req.database)
            cursor = conn.cursor()
            cursor.execute("UPDATE Sessions SET InvalidReason = ? WHERE SessionId = ?", (req.comment, req.session_id))
            if req.comment and req.comment.lower().startswith("fake"):
                cursor.execute("UPDATE Sessions SET Valid = 0 WHERE SessionId = ?", (req.session_id,))
            else:
                cursor.execute("UPDATE Sessions SET Valid = 1 WHERE SessionId = ?", (req.session_id,))
            conn.commit()
            conn.close()
            return {"message": "Comment updated successfully in Sessions"}
        except Exception as fallback_e:
            print(f"Fallback Error in update_call_comment: {fallback_e}")
            raise HTTPException(status_code=500, detail=str(e))

@app.get("/api/calls")
def list_calls(
    database: str = Query(..., min_length=1),
    collection: list[str] | None = Query(default=None),
    location: list[str] | None = Query(default=None),
):
    try:
        conn = get_connection(database)
        cursor = conn.cursor()

        query = """
            SELECT
                DF.ASideLocation AS Location,
                CA.SessionId,
                CA.technology as technology,
                CA.callmode AS callMode,
                CA.callType,
                CA.callDir,
                CA.callStatus AS status,
                DF.CollectionName,
                COALESCE(S.startTime, SB.startTime) AS callStartTimeStamp,
                ROUND(CA.setupTime, 2) AS setupTime,
                (SELECT ROUND(AVG(OptionalWB),2) AS MOS
                    FROM ResultsLQ08Avg 
                    WHERE SessionId = CA.SessionId) AS Avg_mos,
                (ca.callDuration/1000) as callDuration,
                COALESCE (AC.Comment, s.InvalidReason) AS comment,
                DF.ASideFileName,
                POS.Latitude AS latitude,
                POS.Longitude AS longitude,
                S.Valid AS isValid
            FROM CallAnalysis CA
            LEFT JOIN FileList DF ON CA.FileId = DF.FileId
            LEFT JOIN Position POS ON CA.PosId = POS.PosId
            LEFT JOIN Sessions S ON S.SessionId = CA.SessionId
            LEFT JOIN SessionsB SB ON SB.SessionId = CA.SessionId
			LEFT JOIN AnalysisCommentSessionsBridge ACSB ON ACSB.sessionID = CA.SessionId
			LEFT JOIN AnalysisComment AC ON ACSB.commentId = AC.commentID
            WHERE (S.Valid = 1 or S.Valid = 0)
        """

        params: list[object] = []
        selected_collections = [col for col in (collection or []) if col and col.strip()]

        if selected_collections:
            placeholders = ", ".join(["?"] * len(selected_collections))
            query += f" AND DF.CollectionName IN ({placeholders})"
            params.extend(selected_collections)

        selected_locations = [loc for loc in (location or []) if loc and loc.strip()]

        if selected_locations:
            placeholders = ", ".join(["?"] * len(selected_locations))
            query += f" AND DF.ASideLocation IN ({placeholders})"
            params.extend(selected_locations)

        query += " ORDER BY callStartTimeStamp"

        cursor.execute(query, tuple(params))

        columns = [col[0] for col in cursor.description] if cursor.description else []
        rows = cursor.fetchall() if cursor.description else []

        data = []
        for row in rows:
            data.append({columns[idx]: row[idx] for idx in range(len(columns))})

        conn.close()

        return {"rows": data}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.get("/api/locations")
def list_locations(
    database: str = Query(..., min_length=1),
    collection: list[str] | None = Query(default=None),
):
    try:
        conn = get_connection(database)
        cursor = conn.cursor()

        query = """
            SELECT DISTINCT ASideLocation
            FROM FileList
            WHERE ASideLocation IS NOT NULL
        """
        params = []
        
        selected_collections = [col for col in (collection or []) if col and col.strip()]
        if selected_collections:
            placeholders = ", ".join(["?"] * len(selected_collections))
            query += f" AND CollectionName IN ({placeholders})"
            params.extend(selected_collections)
            
        query += " ORDER BY ASideLocation"

        cursor.execute(query, tuple(params))

        rows = cursor.fetchall()
        conn.close()

        return {"locations": [row[0] for row in rows if row[0]]}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@app.get("/api/databases")
def list_databases():
    try:
        return {"databases": get_available_databases()}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@app.post("/api/benchmark")
def run_benchmark(req: QueryRequest):
    try:
        conn = get_connection(req.database)
        cursor = conn.cursor()

        results = []
        total_start = time.time()

        for i, query in enumerate(req.queries):
            q_start = time.time()
            cursor.execute(query)

            columns = [col[0] for col in cursor.description] if cursor.description else []
            rows = cursor.fetchall() if cursor.description else []

            data = []
            for row in rows:
                data.append({columns[idx]: row[idx] for idx in range(len(columns))})

            exec_ms = round((time.time() - q_start) * 1000)

            results.append({
                "id": f"result-{i}",
                "queryLabel": f"Query {i+1}",
                "executionTime": exec_ms,
                "rowsReturned": len(data),
                "columns": columns,
                "data": data
            })

        conn.close()

        total_time = round((time.time() - total_start) * 1000)

        return {
            "results": results,
            "totalTime": total_time
        }

    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.get("/api/collections")
def list_collections(database: str = Query(..., min_length=1)):
    try:
        conn = get_connection(database)
        cursor = conn.cursor()

        cursor.execute("""
            SELECT DISTINCT CollectionName
            FROM filelist
            WHERE CollectionName IS NOT NULL
            ORDER BY CollectionName
        """)

        rows = cursor.fetchall()
        conn.close()

        return {"collections": [row[0] for row in rows if row[0]]}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@app.get("/api/lte_values")
def get_lte_values(
    database: str = Query(..., min_length=1),
    session_id: str = Query(..., min_length=1)
):
    try:
        conn = get_connection(database)
        cursor = conn.cursor()

        query = """
            SELECT lmr.[MsgId]
                  ,lmr.[SessionId]
                  ,lmr.[MsgTime]
                  ,lmr.[PosId]
                  ,lmr.[NetworkId]
                  ,lmr.[EARFCN]
                  ,lmr.[PhyCellId]
                  ,round(lmr.[RSRP], 2) AS [RSRP]
                  ,round(lmr.[RSRQ], 2) AS [RSRQ]
                  ,round(lmr.[SINR0], 2) AS [SINR0]
                  ,round(lmr.[SINR1], 2) AS [SINR1]
                  ,lmr.[LTEServingCellInfoId]
                  ,p.Latitude
                  ,p.Longitude
              FROM [LTEMeasurementReport] lmr
              LEFT JOIN Position p ON p.PosId = lmr.PosId
              WHERE lmr.[SessionId] = ?
              ORDER BY lmr.MsgTime
        """

        cursor.execute(query, (session_id,))

        columns = [col[0] for col in cursor.description] if cursor.description else []
        rows = cursor.fetchall() if cursor.description else []

        data = []
        for row in rows:
            data.append({columns[idx]: row[idx] for idx in range(len(columns))})

        conn.close()

        return {"lteValues": data}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.get("/api/cell_info")
def get_cell_info(
    database: str = Query(..., min_length=1),
    session_id: str = Query(..., min_length=1)
):
    try:
        conn = get_connection(database)
        cursor = conn.cursor()
        cursor.execute("""
            SELECT TOP 1
                dci.eNBId,
                fl.EARFCN,
                fl.PhyCellId
            FROM FactLTERadio fl
            LEFT JOIN DmnCellInformation dci ON fl.DmnIdCellInformation = dci.DmnId
            WHERE fl.SessionId = ?
              AND dci.eNBId IS NOT NULL
            ORDER BY fl.FullDate
        """, (session_id,))
        row = cursor.fetchone()
        conn.close()
        if row:
            return {"eNBId": row[0], "EARFCN": row[1], "PCI": row[2]}
        return {"eNBId": None, "EARFCN": None, "PCI": None}
    except Exception as e:
        return {"eNBId": None, "EARFCN": None, "PCI": None}


@app.get("/api/lte_values_b_side")
def get_lte_values_b_side(
    database: str = Query(..., min_length=1),
    session_id: str = Query(..., min_length=1)
):
    try:
        conn = get_connection(database)
        cursor = conn.cursor()

        query = """
            ;WITH pair_root AS (
                SELECT TOP (1)
                    CASE
                        WHEN CA.Side = 'B' AND CA.SessionIdA IS NOT NULL THEN CA.SessionIdA
                        ELSE CA.SessionId
                    END AS ASessionId
                FROM CallAnalysis CA
                WHERE CA.SessionId = TRY_CONVERT(BIGINT, ?)
                   OR CA.SessionIdA = TRY_CONVERT(BIGINT, ?)
            ),
            b_side AS (
                SELECT TOP (1)
                    CA.SessionId AS BSessionId
                FROM CallAnalysis CA
                INNER JOIN pair_root PR
                    ON CA.SessionIdA = PR.ASessionId
                WHERE CA.Side = 'B'
            )
            SELECT
                L.SessionId,
                L.MsgTime,
                L.PosId,
                L.EARFCN,
                ROUND(L.RSRP, 2) AS RSRP,
                ROUND(L.RSRQ, 2) AS RSRQ,
                ROUND(L.SINR0, 2) AS SINR0,
                ROUND(L.SINR1, 2) AS SINR1,
                P.Latitude,
                P.Longitude
            FROM LTEMeasurementReport L
            INNER JOIN b_side B
                ON L.SessionId = B.BSessionId
            LEFT JOIN Position P
                ON P.PosId = L.PosId
            ORDER BY L.MsgTime
        """

        cursor.execute(query, (session_id, session_id))

        columns = [col[0] for col in cursor.description] if cursor.description else []
        rows = cursor.fetchall() if cursor.description else []

        data = []
        for row in rows:
            data.append({columns[idx]: row[idx] for idx in range(len(columns))})

        conn.close()

        return {"lteValuesBSide": data}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@app.get("/api/markers")
def get_markers(
    database: str = Query(..., min_length=1),
    session_id: str = Query(..., min_length=1)
):
    try:
        conn = get_connection(database)
        cursor = conn.cursor()

        query = """
            SELECT [markerId]
                  ,[SessionId]
                  ,[MsgTime]
                  ,[PosId]
                  ,[NetworkId]
                  ,[MarkerText]
              FROM [Markers]
              WHERE [SessionId] = ?
              ORDER BY MsgTime
        """

        cursor.execute(query, (session_id,))

        columns = [col[0] for col in cursor.description] if cursor.description else []
        rows = cursor.fetchall() if cursor.description else []

        data = []
        for row in rows:
            data.append({columns[idx]: row[idx] for idx in range(len(columns))})

        conn.close()

        return {"markers": data}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@app.get("/api/gsm_values")
def get_gsm_values(
    database: str = Query(..., min_length=1),
    session_id: str = Query(..., min_length=1)
):
    try:
        conn = get_connection(database)
        cursor = conn.cursor()

        query = """
            SELECT g.[MsgId]
                  ,g.[SessionId]
                  ,g.[MsgTime]
                  ,g.[PosId]
                  ,g.[NetworkId]
                  ,g.[RxLevSub]
                  ,g.[RxQualSub]
                  ,p.Latitude
                  ,p.Longitude
              FROM [GSMMeasReport] g
              LEFT JOIN Position p ON p.PosId = g.PosId
              WHERE g.[SessionId] = ?
              ORDER BY g.MsgTime
        """

        cursor.execute(query, (session_id,))

        columns = [col[0] for col in cursor.description] if cursor.description else []
        rows = cursor.fetchall() if cursor.description else []

        data = []
        for row in rows:
            data.append({columns[idx]: row[idx] for idx in range(len(columns))})

        conn.close()

        return {"gsmValues": data}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@app.get("/api/gsm_values_b_side")
def get_gsm_values_b_side(
    database: str = Query(..., min_length=1),
    session_id: str = Query(..., min_length=1)
):
    try:
        conn = get_connection(database)
        cursor = conn.cursor()

        query = """
            ;WITH pair_root AS (
                SELECT TOP (1)
                    CASE
                        WHEN CA.Side = 'B' AND CA.SessionIdA IS NOT NULL THEN CA.SessionIdA
                        ELSE CA.SessionId
                    END AS ASessionId
                FROM CallAnalysis CA
                WHERE CA.SessionId = TRY_CONVERT(BIGINT, ?)
                   OR CA.SessionIdA = TRY_CONVERT(BIGINT, ?)
            ),
            b_side AS (
                SELECT TOP (1)
                    CA.SessionId AS BSessionId
                FROM CallAnalysis CA
                INNER JOIN pair_root PR
                    ON CA.SessionIdA = PR.ASessionId
                WHERE CA.Side = 'B'
            )
            SELECT
                G.[MsgId],
                G.[SessionId],
                G.[MsgTime],
                G.[PosId],
                G.[NetworkId],
                G.[RxLevSub],
                G.[RxQualSub]
            FROM [GSMMeasReport] G
            INNER JOIN b_side B
                ON G.SessionId = B.BSessionId
            ORDER BY G.MsgTime
        """

        cursor.execute(query, (session_id, session_id))

        columns = [col[0] for col in cursor.description] if cursor.description else []
        rows = cursor.fetchall() if cursor.description else []

        data = []
        for row in rows:
            data.append({columns[idx]: row[idx] for idx in range(len(columns))})

        conn.close()

        return {"gsmValuesBSide": data}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@app.get("/api/mos_values")
def get_mos_values(
    database: str = Query(..., min_length=1),
    session_id: str = Query(..., min_length=1)
):
    try:
        conn = get_connection(database)
        cursor = conn.cursor()

        query = """
            SELECT OptionalWB
              FROM [ResultsLQ08Avg]
              WHERE [SessionId] = ?
              ORDER BY MsgId
        """

        cursor.execute(query, (session_id,))

        columns = [col[0] for col in cursor.description] if cursor.description else []
        rows = cursor.fetchall() if cursor.description else []

        data = []
        for row in rows:
            data.append({columns[idx]: row[idx] for idx in range(len(columns))})

        conn.close()

        return {"mosValues": data}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@app.get("/api/results_kpi")
def get_results_kpi(
    database: str = Query(..., min_length=1),
    session_id: str | None = Query(default=None)
):
    try:
        conn = get_connection(database)
        cursor = conn.cursor()

        query = """
            SELECT [MsgId]
                  ,[SessionId]
                  ,[TestId]
                  ,[KPIId]
                  ,[StartTime]
                  ,[EndTime]
                  ,[ErrorCode]
                  ,[Counter]
                  ,[Value1]
                  ,[Value2]
                  ,[Value3]
                  ,[Value4]
                  ,[Value5]
              FROM [ResultsKPI]
        """
        
        params = []
        if session_id:
            query += " WHERE [SessionId] = ?"
            params.append(session_id)

        query += " ORDER BY [MsgId]"

        cursor.execute(query, tuple(params))

        columns = [col[0] for col in cursor.description] if cursor.description else []
        rows = cursor.fetchall() if cursor.description else []

        data = []
        for row in rows:
            data.append({columns[idx]: row[idx] for idx in range(len(columns))})

        conn.close()

        return {"kpiValues": data}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@app.get("/api/tracelog_values")
def get_tracelog_values(
    database: str = Query(..., min_length=1),
    session_id: str | None = Query(default=None)
):
    try:
        conn = get_connection(database)
        cursor = conn.cursor()

        query = """
           ;WITH pair_root AS (
                SELECT TOP (1)
                    CASE
                        WHEN CA.Side = 'B' AND CA.SessionIdA IS NOT NULL THEN CA.SessionIdA
                        ELSE CA.SessionId
                    END AS ASessionId
                FROM CallAnalysis CA
                WHERE CA.SessionId = TRY_CONVERT(BIGINT, ?)
                   OR CA.SessionIdA = TRY_CONVERT(BIGINT, ?)
            ),
            sessions_to_include AS (
                SELECT PR.ASessionId AS SessionId
                FROM pair_root PR
                WHERE PR.ASessionId IS NOT NULL

                UNION

                SELECT CA.SessionId AS SessionId
                FROM CallAnalysis CA
                INNER JOIN pair_root PR
                    ON CA.SessionIdA = PR.ASessionId
                WHERE CA.Side = 'B'

                UNION

                SELECT TRY_CONVERT(BIGINT, ?)
                WHERE TRY_CONVERT(BIGINT, ?) IS NOT NULL
            )
            SELECT
                TL.[FactId],
                TL.[FullDate],
                TL.[Side],
                TL.[SessionId],
                TL.[Info]
            FROM [dbo].[FactSystemTraceLog] TL
            INNER JOIN sessions_to_include SI
                ON TL.[SessionId] = SI.[SessionId]
            ORDER BY TL.[FullDate], TL.[FactId]
        """

        cursor.execute(query, (session_id, session_id, session_id, session_id))

        columns = [col[0] for col in cursor.description] if cursor.description else []
        rows = cursor.fetchall() if cursor.description else []

        data = []
        for row in rows:
            data.append({columns[idx]: row[idx] for idx in range(len(columns))})

        conn.close()

        return {"tracelogValues": data}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))



@app.get("/api/call_side_comparison")
def get_call_side_comparison(
    database: str = Query(..., min_length=1),
    session_id: str = Query(..., min_length=1)
):
    try:
        conn = get_connection(database)
        cursor = conn.cursor()

        query = """
            WITH root_session AS (
                SELECT TOP (1)
                    CASE
                        WHEN CA.Side = 'B' AND CA.SessionIdA IS NOT NULL THEN CA.SessionIdA
                        ELSE CA.SessionId
                    END AS ASessionId
                FROM CallAnalysis CA
                WHERE CA.SessionId = TRY_CONVERT(BIGINT, ?)
                   OR CA.SessionIdA = TRY_CONVERT(BIGINT, ?)
            )
            SELECT
                CA.Side,
                CA.callStatus,
                CA.code,
                CA.codeDescription,
                COUNT(*) AS calls
            FROM CallAnalysis CA
            CROSS JOIN root_session RS
            WHERE
                (CA.Side = 'A' AND CA.SessionId = RS.ASessionId)
                OR
                (CA.Side = 'B' AND CA.SessionIdA = RS.ASessionId)
            GROUP BY
                CA.Side,
                CA.callStatus,
                CA.code,
                CA.codeDescription
            ORDER BY
                calls DESC
        """

        cursor.execute(query, (session_id, session_id))

        columns = [col[0] for col in cursor.description] if cursor.description else []
        rows = cursor.fetchall() if cursor.description else []

        data = []
        for row in rows:
            data.append({columns[idx]: row[idx] for idx in range(len(columns))})

        conn.close()

        return {"comparison": data}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.get("/api/antennas")
def get_antennas(
    freq: list[int] | None = Query(default=None),
    vendor: list[str] | None = Query(default=None),
    status: str | None = Query(default=None),
):
    """Return all 4G antenna sectors from geo4g.xlsx.
    Optional filters: freq (e.g. 1800, 2100), vendor, status (ACTIVATED/DEACTIVATED).
    First call loads & caches the file (~3-5s); subsequent calls are instant.
    """
    try:
        antennas = _load_antennas()
        result = antennas

        if freq:
            freq_set = set(freq)
            result = [a for a in result if a["freq"] in freq_set]

        if vendor:
            vendor_set = {v.lower() for v in vendor}
            result = [a for a in result if (a["vendor"] or "").lower() in vendor_set]

        if status:
            result = [a for a in result if (a["status"] or "").upper() == status.upper()]

        return {"antennas": result, "total": len(result)}
    except FileNotFoundError:
        raise HTTPException(status_code=404, detail="geo4g.xlsx not found")
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
